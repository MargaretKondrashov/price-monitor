import asyncio
import json
import os
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

for k in ("HTTPS_PROXY", "HTTP_PROXY", "https_proxy", "http_proxy", "ALL_PROXY", "all_proxy"):
    os.environ.pop(k, None)

from monitor.extract import extract_page, parse_price  # noqa: E402
from monitor.models import extract_model, find_exact, find_similar, normalize  # noqa: E402
from monitor.pipeline import run, save  # noqa: E402
from monitor.report import build_email_html, build_excel  # noqa: E402

from . import fake_sites  # noqa: E402


# ---------------- unit: model matching ----------------

def test_extract_model_from_traklin_titles():
    assert extract_model('מסך "75 SAMSUNG SMART 4K דגם UE75U8000FUXIL סמסונג') == "UE75U8000FUXIL"
    assert extract_model("מקרר 4 דלתות LG דגם GR-728B נירוסטה") == "GR-728B"
    assert extract_model("", "/x/מקרר_haier_דגם_hrf-825fb_זכוכית") == "HRF-825FB"
    assert extract_model("תנור בילט אין יוקרתי") is None
    assert extract_model("מזגן 12000BTU 4K") is None


def test_exact_match_tolerates_separators_but_not_longer_models():
    assert find_exact("GR-728B", "מקרר LG GR728B נירוסטה")
    assert find_exact("GR-728B", "מק\"ט: 510.2021.gr-728b")
    assert not find_exact("GR-728B", "מקרר LG GR-728BS")
    assert not find_exact("GR-728B", "מקרר LG GR-728S")


def test_similar_model_rules():
    assert find_similar("UE75U8000FUXIL", "Samsung UE75U8000HUXIL 75") == "UE75U8000HUXIL"
    # different size must never be "similar"
    assert find_similar("UE75U8000FUXIL", "Samsung UE65U8000FUXIL") is None
    # different color variant (one char out of six) stays below threshold
    assert find_similar("GR-728B", "LG GR-728S") is None


def test_parse_price():
    assert parse_price("3,049 ₪") == 3049
    assert parse_price("₪ 1,299.90") == 1299.9
    assert parse_price("0") is None
    assert parse_price("לא") is None


def test_heuristic_ignores_related_products_and_reads_old_price():
    html = """<main><h1>X</h1><div class="price-box"><del>1,999 ₪</del><span class="final-price">1,499 ₪</span></div>
    <div class="related"><span class="price">10 ₪</span></div></main>"""
    info = extract_page(html)
    assert info.price == 1499 and info.old_price == 1999 and info.discount_pct == 25.0


def test_model_in_carousel_or_menu_does_not_confirm_a_different_product():
    html = """<header>GR-728B</header><main><h1>מקרר LG GR-728S</h1><span class="price">8,990 ₪</span>
    <div class="related-products"><a>מקרר LG GR-728B</a></div></main><footer>GR-728B</footer>"""
    info = extract_page(html)
    assert not find_exact("GR-728B", f"{info.title} {info.text}")
    assert find_exact("GR-728S", info.text) and info.price == 8990


# ---------------- end to end against fake stores ----------------

@pytest.fixture(scope="module")
def world():
    return fake_sites.start(8765)


def make_cfg(urls):
    return {
        "run": {"max_runtime_minutes": 5, "recheck_not_found_days": 7, "site_probe_products": 3, "http_retries": 0},
        "matching": {"allow_similar_models": True, "similarity_threshold": 0.86},
        "report": {"min_change_percent": 0.5},
        "traklin": {"name": "טרקלין חשמל", "base_url": urls["traklin"], "sitemap_pages": ["/sitemap.aspx"],
                    "category_seeds": ["/נקיון_וכביסה"], "concurrency": 4, "delay_seconds": 0},
        "stores": [
            {"key": "payngo", "name": "מחסני חשמל", "type": "generic", "base_url": urls["magento"],
             "search_templates": ["/search?q={q}"]},           # wrong guess: must self-discover the real one
            {"key": "ksp", "name": "KSP", "type": "ksp", "base_url": urls["ksp"],
             "api_url": urls["ksp"] + "/m_action/api", "item_url": urls["ksp"] + "/web/item/"},
            {"key": "bigelectric", "name": "ביג אלקטריק", "type": "generic", "base_url": urls["woo"]},
            {"key": "bug", "name": "באג", "type": "generic", "base_url": urls["custom"]},
            {"key": "blocked", "name": "חסום", "type": "generic", "base_url": urls["blocked"],
             "search_templates": ["/search?q={q}"]},
        ],
        "store_defaults": {"concurrency": 3, "delay_seconds": 0},
    }


def test_full_pipeline_two_days(world, tmp_path):
    cfg = make_cfg(world)
    data, dbg, out = tmp_path / "data", tmp_path / "debug", tmp_path / "out"

    # ---- day 1 ----
    r1 = asyncio.run(run(cfg, data, dbg))
    models = {normalize(p.model): p for p in r1.products}
    assert set(models) == {"UE75U8000FUXIL", "GR728B", "WGB244A0IL", "HRF825FB"}, models.keys()
    assert r1.health["traklin"]["no_model"] == 1
    assert models["GR728B"].price == 8990 and models["GR728B"].old_price == 9990
    assert models["WGB244A0IL"].price == 3490        # found only on category page 2
    assert r1.first_run

    o = r1.offers
    # Magento: self-discovered search form, exact matches, old price, carousel ignored
    assert o["payngo"]["GR728B"].price == 8790 and o["payngo"]["GR728B"].old_price == 9490
    assert o["payngo"]["GR728B"].match == "exact"
    assert o["payngo"]["UE75U8000FUXIL"].price == 3349
    # KSP: campaign discount applied, labels kept
    k = o["ksp"]["UE75U8000FUXIL"]
    assert (k.price, k.old_price) == (3190, 3390) and any("הנחת חג" in p for p in k.promos)
    assert o["ksp"]["GR728B"].status == "not_found"
    # Woo: store API, regular price as old price, GR-728S must NOT be matched to GR-728B
    w = o["bigelectric"]["WGB244A0IL"]
    assert (w.price, w.old_price, w.match) == (3290, 3690, "exact")
    assert o["bigelectric"]["GR728B"].status == "not_found"
    # Custom: search redirects to product; HUXIL accepted only as "similar"
    b = o["bug"]["UE75U8000FUXIL"]
    assert (b.price, b.match, b.matched_model) == (3049, "similar", "UE75U8000HUXIL")
    assert o["bug"]["GR728B"].price == 8890 and o["bug"]["GR728B"].match == "exact"
    # Blocked site flagged, not crashing
    hb = next(h for h in r1.health["stores"] if h["site"] == "blocked")
    assert hb["found"] == 0 and hb["broken"]

    save(r1, data)
    assert (data / "latest.json").exists() and (data / "url_cache.json").exists()
    cache = json.loads((data / "url_cache.json").read_text("utf-8"))
    assert not cache["urls"].get("blocked"), "a blocked site must not be remembered as 'not found'"
    assert cache["urls"]["bigelectric"]["GR728B"]["url"] is None    # answered, genuinely not there

    # ---- day 2: prices move ----
    fake_sites.WORLD["magento"]["337350"] = ("מקרר 4 דלתות 665 ליטר LG GR-728B התקן שבת", 8490, 9490, "GR-728B")
    fake_sites.WORLD["traklin"]["נקיון_וכביסה/מכונות_כביסה/מכונת_כביסה_bosch_דגם_wgb244a0il"] = (
        "מכונת כביסה BOSCH 9 ק\"ג דגם WGB244A0IL", 3390, 3690)
    fake_sites.WORLD["custom"]["9577452452"] = ("מקרר 4 דלתות אינוורטר - 665 ליטר LG GR-728B נירוסטה מושחרת", 9190)
    r2 = asyncio.run(run(cfg, data, dbg))
    assert not r2.first_run and r2.prev_date
    moves = {(c.site, normalize(c.model)): c for c in r2.changes if c.kind in ("up", "down")}
    assert moves[("payngo", "GR728B")].kind == "down" and moves[("payngo", "GR728B")].pct == -3.4
    assert moves[("bug", "GR728B")].kind == "up"
    assert moves[("traklin", "WGB244A0IL")].new == 3390
    assert any(c.kind == "new_promo" and c.site == "traklin" for c in r2.changes)
    save(r2, data)
    assert "GR-728B" in (data / "price_changes.csv").read_text("utf-8-sig")

    # ---- report ----
    xlsx = build_excel(r2, out / "r.xlsx")
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["סיכום", "השוואה מלאה", "שינויי מחיר", "אירועים", "מבצעים", "מתחרים זולים יותר"]
    assert wb["השוואה מלאה"].sheet_view.rightToLeft
    cheaper = [row[0].value for row in wb["מתחרים זולים יותר"].iter_rows(min_row=2) if row[0].value]
    assert any("GR-728B" in t for t in cheaper)        # payngo 8490 < traklin 8990
    subject, body = build_email_html(r2)
    assert "שינויי מחיר" in subject and 'dir="rtl"' in body and "⚠ לא נקרא היום" in body
    (out / "email.html").write_text(body, "utf-8")
    shutil.copy(xlsx, Path(os.environ.get("REPORT_COPY_DIR", out)) / "test_report.xlsx") \
        if os.environ.get("REPORT_COPY_DIR") else None
