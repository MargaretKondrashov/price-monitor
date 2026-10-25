"""Excel report (Hebrew, right-to-left) and the HTML email body."""
from __future__ import annotations

import datetime as dt
import html
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import normalize
from .pipeline import RunResult

# ---------- palette ----------
NAVY = "1F3A5F"
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor="DCE6F1")
TRAKLIN_FILL = PatternFill("solid", fgColor="FFF4D6")
CHEAPER_FILL = PatternFill("solid", fgColor="FADBD8")    # competitor cheaper than Traklin
DEARER_FILL = PatternFill("solid", fgColor="D5F5E3")     # competitor more expensive
BEST_FILL = PatternFill("solid", fgColor="FCF3CF")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
GREY = Font(color="808080", italic=True)
RED = Font(color="C0392B", bold=True)
GREEN = Font(color="1E8449", bold=True)
LINK = Font(color="1F5FA8", underline="single")
THIN = Side(style="thin", color="D0D7DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NIS = '#,##0" ₪"'
PCT = '0.0"%"'
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT_WRAP = Alignment(horizontal="right", vertical="top", wrap_text=True)

KIND_HE = {"up": "עלייה", "down": "ירידה", "new_promo": "מבצע חדש", "promo_ended": "מבצע הסתיים",
           "back_in_stock": "חזר למלאי", "out_of_stock": "אזל מהמלאי"}


def _fmt_date(iso: str | None) -> str:
    if not iso:
        return ""
    d = dt.date.fromisoformat(iso)
    return d.strftime("%d.%m.%Y")


def _sheet(wb: Workbook, title: str, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.sheet_view.rightToLeft = True
    return ws


def _header_row(ws, row: int, values: list[str], fill=HEAD_FILL, font=WHITE_BOLD):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.fill, c.font, c.alignment, c.border = fill, font, CENTER, BORDER


def _widths(ws, widths: list[int]):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _change_map(result: RunResult) -> dict:
    """(site, normalized model) -> Change of kind up/down"""
    return {(c.site, normalize(c.model)): c for c in result.changes if c.kind in ("up", "down")}


def _promo_text(promos, old, price) -> str:
    parts = []
    if old and price and old > price * 1.01:
        parts.append(f"במקום {old:,.0f} ₪ (‎-{(1 - price / old) * 100:.0f}%)")
    parts += list(promos or [])[:2]
    return " | ".join(parts)


def stats_per_store(result: RunResult) -> list[dict]:
    rows = []
    tk = {normalize(p.model): p for p in result.products}
    for key, name in result.stores:
        offs = result.offers.get(key, {})
        found = [o for o in offs.values() if o.status in ("ok", "no_price")]
        cheaper = dearer = same = 0
        gaps = []
        for m, o in offs.items():
            p = tk.get(m)
            if not (o.price and p and p.price):
                continue
            g = (o.price - p.price) / p.price * 100
            gaps.append(g)
            if g < -0.5:
                cheaper += 1
            elif g > 0.5:
                dearer += 1
            else:
                same += 1
        h = next((x for x in result.health.get("stores", []) if x["site"] == key), {})
        rows.append({
            "key": key, "name": name, "found": len(found), "cheaper": cheaper, "dearer": dearer, "same": same,
            "avg_gap": round(sum(gaps) / len(gaps), 1) if gaps else None,
            "promos": sum(1 for o in found if o.promos or o.discount_pct),
            "ok": not h.get("broken") and len(found) > 0,
            "notes": h.get("notes", []),
        })
    return rows


def cheaper_elsewhere(result: RunResult) -> list[dict]:
    out = []
    names = dict(result.stores)
    for p in result.products:
        if not p.price:
            continue
        m = normalize(p.model)
        best = None
        for key, offs in result.offers.items():
            o = offs.get(m)
            if o and o.price and o.match == "exact" and (best is None or o.price < best.price):
                best = o
        if best and best.price < p.price * 0.995:
            out.append({"product": p, "offer": best, "store": names[best.site],
                        "gap": p.price - best.price, "gap_pct": (p.price - best.price) / p.price * 100})
    out.sort(key=lambda r: r["gap"], reverse=True)
    return out


def all_promos(result: RunResult, traklin_name: str) -> list[dict]:
    rows = []
    names = dict(result.stores)
    for p in result.products:
        if p.promos or p.discount_pct:
            rows.append({"store": traklin_name, "title": p.title, "model": p.model, "price": p.price,
                         "old": p.old_price, "disc": p.discount_pct, "text": " | ".join(p.promos[:3]), "url": p.url})
        m = normalize(p.model)
        for key, offs in result.offers.items():
            o = offs.get(m)
            if o and o.status == "ok" and (o.promos or o.discount_pct):
                rows.append({"store": names[key], "title": p.title, "model": p.model, "price": o.price,
                             "old": o.old_price, "disc": o.discount_pct, "text": " | ".join(o.promos[:3]), "url": o.url})
    rows.sort(key=lambda r: (r["disc"] or 0), reverse=True)
    return rows


# =====================================================================
def build_excel(result: RunResult, path: Path, traklin_name: str = "טרקלין חשמל") -> Path:
    wb = Workbook()
    stores = result.stores
    changes = _change_map(result)
    store_stats = stats_per_store(result)

    # ---------- 1. summary ----------
    ws = _sheet(wb, "סיכום", first=True)
    ws["A1"] = f"דוח השוואת מחירים – {traklin_name} מול המתחרים"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = (f"תאריך הריצה: {_fmt_date(result.date)}" +
                (f"   |   השוואה לריצה מתאריך {_fmt_date(result.prev_date)}" if result.prev_date else
                 "   |   ריצה ראשונה – שינויים יופיעו החל מהריצה הבאה"))
    ws["A2"].font = GREY
    if result.partial:
        ws["A3"] = "⚠ הריצה נעצרה לפני הסוף בגלל מגבלת זמן – חלק מהמוצרים לא נבדקו היום."
        ws["A3"].font = RED
    th = result.health.get("traklin", {})
    kpis = [("מוצרים בטרקלין (עם דגם)", th.get("products", len(result.products))),
            ("השוואות שנמצאו אצל מתחרים", sum(s["found"] for s in store_stats)),
            ("שינויי מחיר מאז הריצה הקודמת", sum(1 for c in result.changes if c.kind in ("up", "down"))),
            ("מבצעים פעילים (כל האתרים)", len(all_promos(result, traklin_name))),
            ("מוצרים שזולים יותר אצל מתחרה", len(cheaper_elsewhere(result)))]
    r = 5
    for label, val in kpis:
        ws.cell(row=r, column=1, value=label).font = BOLD
        c = ws.cell(row=r, column=2, value=val)
        c.font = Font(bold=True, size=13, color=NAVY)
        r += 1
    r += 1
    _header_row(ws, r, ["חנות", "מצב", "מוצרים שנמצאו", "זול מטרקלין", "יקר מטרקלין", "אותו מחיר",
                        "הפרש ממוצע מטרקלין", "מבצעים", "הערות"])
    for s in store_stats:
        r += 1
        vals = [s["name"], "✓ תקין" if s["ok"] else "⚠ לא נקרא", s["found"], s["cheaper"], s["dearer"], s["same"],
                s["avg_gap"], s["promos"], "" if s["ok"] else " ".join(s["notes"][-1:])]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BORDER
            c.alignment = RIGHT_WRAP if i == 9 else CENTER
        ws.cell(row=r, column=2).font = GREEN if s["ok"] else RED
        ws.cell(row=r, column=7).number_format = '+0.0"%";-0.0"%";0"%"'
    r += 2
    ws.cell(row=r, column=1, value="מקרא צבעים בגיליון 'השוואה מלאה':").font = BOLD
    for fill, txt in ((CHEAPER_FILL, "המתחרה זול יותר מטרקלין"), (DEARER_FILL, "המתחרה יקר יותר מטרקלין"),
                      (BEST_FILL, "המחיר הזול ביותר בשוק"), (None, "טקסט אפור נטוי = דגם דומה ולא זהה – כדאי לבדוק")):
        r += 1
        c = ws.cell(row=r, column=1, value=txt)
        if fill:
            c.fill = fill
        else:
            c.font = GREY
    _widths(ws, [34, 12, 15, 13, 13, 12, 18, 10, 60])

    # ---------- 2. full comparison ----------
    ws = _sheet(wb, "השוואה מלאה")
    base_cols = ["קטגוריה", "מוצר", "דגם", "מחיר", "שינוי", "מבצע / מחיר מחוק"]
    per_store = ["מחיר", "הפרש מטרקלין", "שינוי", "מבצע / הערה"]
    tail = ["הכי זול בשוק", "מחיר הכי זול"]
    # two-row header
    ws.cell(row=1, column=1, value="פרטי מוצר")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    ws.cell(row=1, column=4, value=traklin_name)
    ws.merge_cells(start_row=1, start_column=4, end_row=1, end_column=6)
    col = 7
    store_col = {}
    for key, name in stores:
        store_col[key] = col
        ws.cell(row=1, column=col, value=name)
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(per_store) - 1)
        col += len(per_store)
    ws.cell(row=1, column=col, value="סיכום")
    ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)
    total_cols = col + 1
    for i in range(1, total_cols + 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font, c.alignment, c.border = HEAD_FILL, WHITE_BOLD, CENTER, BORDER
    sub = base_cols + per_store * len(stores) + tail
    _header_row(ws, 2, sub, fill=SUB_FILL, font=BOLD)

    def change_cell(cell, ch):
        if ch and ch.pct is not None:
            cell.value = f"{'▲' if ch.pct > 0 else '▼'} {abs(ch.pct):.1f}%"
            cell.font = RED if ch.pct > 0 else GREEN
            cell.alignment = CENTER

    names = dict(stores)
    products = sorted(result.products, key=lambda p: (p.category, p.title))
    for ri, p in enumerate(products, start=3):
        m = normalize(p.model)
        ws.cell(row=ri, column=1, value=p.category)
        tc = ws.cell(row=ri, column=2, value=p.title)
        if p.url:
            tc.hyperlink = p.url
            tc.font = LINK
        ws.cell(row=ri, column=3, value=p.model)
        pc = ws.cell(row=ri, column=4, value=p.price)
        pc.number_format, pc.fill, pc.font = NIS, TRAKLIN_FILL, BOLD
        change_cell(ws.cell(row=ri, column=5), changes.get(("traklin", m)))
        ws.cell(row=ri, column=6, value=_promo_text(p.promos, p.old_price, p.price))

        best_price, best_store = p.price, traklin_name if p.price else None
        price_cells = []
        for key, _ in stores:
            c0 = store_col[key]
            o = result.offers.get(key, {}).get(m)
            if not o or o.status not in ("ok", "no_price"):
                note = {"not_found": "לא נמצא", "site_error": "האתר לא נקרא", "skipped": "לא נבדק"}.get(
                    o.status if o else "not_found", "")
                nc = ws.cell(row=ri, column=c0 + 3, value=note)
                nc.font = GREY
                continue
            pc = ws.cell(row=ri, column=c0, value=o.price)
            pc.number_format = NIS
            if o.url:
                pc.hyperlink = o.url
            similar = o.match == "similar"
            pc.font = GREY if similar else (LINK if o.url else BOLD)
            if o.price and p.price:
                gap = o.price - p.price
                gc = ws.cell(row=ri, column=c0 + 1, value=gap)
                gc.number_format = '+#,##0" ₪";-#,##0" ₪";0'
                if not similar:
                    if gap < -0.5:
                        gc.fill = CHEAPER_FILL
                    elif gap > 0.5:
                        gc.fill = DEARER_FILL
                if not similar and (best_price is None or o.price < best_price):
                    best_price, best_store = o.price, names[key]
                price_cells.append((o.price, pc, similar))
            change_cell(ws.cell(row=ri, column=c0 + 2), changes.get((key, m)))
            note = _promo_text(o.promos, o.old_price, o.price)
            if similar:
                note = f"דגם דומה: {o.matched_model}" + (f" | {note}" if note else "")
            if o.in_stock is False:
                note = "אזל מהמלאי" + (f" | {note}" if note else "")
            if o.status == "no_price":
                note = "מחיר לא זוהה" + (f" | {note}" if note else "")
            ws.cell(row=ri, column=c0 + 3, value=note)
        if best_price:
            if p.price and best_price >= p.price - 0.5:
                ws.cell(row=ri, column=4).fill = BEST_FILL
            for price, cell, similar in price_cells:
                if not similar and abs(price - best_price) < 0.5:
                    cell.fill = BEST_FILL
        ws.cell(row=ri, column=total_cols - 1, value=best_store)
        bc = ws.cell(row=ri, column=total_cols, value=best_price)
        bc.number_format = NIS
        for i in range(1, total_cols + 1):
            ws.cell(row=ri, column=i).border = BORDER

    widths = [18, 46, 16, 11, 9, 26] + [11, 12, 9, 24] * len(stores) + [14, 12]
    _widths(ws, widths)
    ws.freeze_panes = "D3"
    ws.auto_filter.ref = f"A2:{get_column_letter(total_cols)}{max(3, len(products) + 2)}"
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 30

    # ---------- 3. price changes ----------
    ws = _sheet(wb, "שינויי מחיר")
    _header_row(ws, 1, ["חנות", "מוצר", "דגם", "מחיר קודם", "מחיר חדש", "שינוי %", "סוג", "קישור"])
    rows = sorted([c for c in result.changes if c.kind in ("up", "down")], key=lambda c: abs(c.pct or 0), reverse=True)
    for ri, c in enumerate(rows, start=2):
        vals = [c.site_name, c.title, c.model, c.old, c.new, c.pct, KIND_HE[c.kind], "לעמוד המוצר" if c.url else ""]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=i, value=v)
            cell.border = BORDER
        ws.cell(row=ri, column=4).number_format = NIS
        ws.cell(row=ri, column=5).number_format = NIS
        ws.cell(row=ri, column=6).number_format = '+0.0"%";-0.0"%"'
        ws.cell(row=ri, column=6).font = RED if (c.pct or 0) > 0 else GREEN
        if c.url:
            ws.cell(row=ri, column=8).hyperlink = c.url
            ws.cell(row=ri, column=8).font = LINK
    if not rows:
        ws.cell(row=2, column=1, value="ריצה ראשונה – אין עדיין נתונים להשוואה" if result.first_run
                else "לא היו שינויי מחיר מאז הריצה הקודמת").font = GREY
    _widths(ws, [14, 50, 16, 12, 12, 10, 10, 14])
    ws.freeze_panes = "A2"

    # ---------- 4. other events (promos started/ended, stock) ----------
    ws = _sheet(wb, "אירועים")
    _header_row(ws, 1, ["חנות", "מוצר", "דגם", "אירוע", "מחיר", "קישור"])
    ev = [c for c in result.changes if c.kind not in ("up", "down")]
    for ri, c in enumerate(ev, start=2):
        vals = [c.site_name, c.title, c.model, KIND_HE[c.kind], c.new, "לעמוד המוצר" if c.url else ""]
        for i, v in enumerate(vals, 1):
            ws.cell(row=ri, column=i, value=v).border = BORDER
        ws.cell(row=ri, column=5).number_format = NIS
        if c.url:
            ws.cell(row=ri, column=6).hyperlink = c.url
            ws.cell(row=ri, column=6).font = LINK
    if not ev:
        ws.cell(row=2, column=1, value="אין אירועים חדשים").font = GREY
    _widths(ws, [14, 50, 16, 14, 12, 14])
    ws.freeze_panes = "A2"

    # ---------- 5. promotions ----------
    ws = _sheet(wb, "מבצעים")
    _header_row(ws, 1, ["חנות", "מוצר", "דגם", "מחיר", "מחיר מחוק", "הנחה %", "פרטי המבצע", "קישור"])
    promos = all_promos(result, traklin_name)
    for ri, pr in enumerate(promos, start=2):
        vals = [pr["store"], pr["title"], pr["model"], pr["price"], pr["old"], pr["disc"], pr["text"],
                "לעמוד המוצר" if pr["url"] else ""]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=i, value=v)
            cell.border = BORDER
        ws.cell(row=ri, column=4).number_format = NIS
        ws.cell(row=ri, column=5).number_format = NIS
        ws.cell(row=ri, column=6).number_format = PCT
        ws.cell(row=ri, column=7).alignment = RIGHT_WRAP
        if pr["url"]:
            ws.cell(row=ri, column=8).hyperlink = pr["url"]
            ws.cell(row=ri, column=8).font = LINK
    if not promos:
        ws.cell(row=2, column=1, value="לא זוהו מבצעים").font = GREY
    _widths(ws, [14, 46, 16, 11, 11, 9, 60, 14])
    ws.freeze_panes = "A2"

    # ---------- 6. where competitors are cheaper ----------
    ws = _sheet(wb, "מתחרים זולים יותר")
    _header_row(ws, 1, ["מוצר", "דגם", "מחיר בטרקלין", "החנות הזולה", "המחיר שם", "פער ₪", "פער %", "קישור"])
    ce = cheaper_elsewhere(result)
    for ri, row in enumerate(ce, start=2):
        p, o = row["product"], row["offer"]
        vals = [p.title, p.model, p.price, row["store"], o.price, row["gap"], round(row["gap_pct"], 1),
                "לעמוד המוצר" if o.url else ""]
        for i, v in enumerate(vals, 1):
            ws.cell(row=ri, column=i, value=v).border = BORDER
        for col_i in (3, 5, 6):
            ws.cell(row=ri, column=col_i).number_format = NIS
        ws.cell(row=ri, column=7).number_format = PCT
        ws.cell(row=ri, column=6).fill = CHEAPER_FILL
        if o.url:
            ws.cell(row=ri, column=8).hyperlink = o.url
            ws.cell(row=ri, column=8).font = LINK
    if not ce:
        ws.cell(row=2, column=1, value="לא נמצאו מוצרים שזולים יותר אצל המתחרים").font = GREY
    _widths(ws, [50, 16, 13, 14, 12, 11, 9, 14])
    ws.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# =====================================================================
def _e(s) -> str:
    return html.escape(str(s if s is not None else ""))


def _money(v) -> str:
    return f"{v:,.0f} ₪" if v else "—"


def build_email_html(result: RunResult, traklin_name: str = "טרקלין חשמל", top: int = 15) -> tuple[str, str]:
    """Returns (subject, html). Inline styles only – email clients drop <style> blocks."""
    stats = stats_per_store(result)
    moves = [c for c in result.changes if c.kind in ("up", "down")]
    downs = sorted([c for c in moves if c.kind == "down"], key=lambda c: c.pct or 0)[:top]
    ups = sorted([c for c in moves if c.kind == "up"], key=lambda c: -(c.pct or 0))[:top]
    new_promos = [c for c in result.changes if c.kind == "new_promo"][:top]
    cheaper = cheaper_elsewhere(result)[:top]
    promos_total = len(all_promos(result, traklin_name))
    found_total = sum(s["found"] for s in stats)
    date_he = _fmt_date(result.date)
    subject = f"דוח מחירים {date_he}: {len(moves)} שינויי מחיר, {len(cheaper_elsewhere(result))} מוצרים זולים יותר אצל מתחרים"

    td = 'style="padding:6px 8px;border-bottom:1px solid #e5e7eb;text-align:right;font-size:13px"'
    th = 'style="padding:6px 8px;background:#1f3a5f;color:#fff;text-align:right;font-size:13px"'

    def table(headers, rows):
        h = "".join(f"<th {th}>{_e(x)}</th>" for x in headers)
        b = "".join("<tr>" + "".join(f"<td {td}>{cell}</td>" for cell in r) + "</tr>" for r in rows)
        return f'<table cellspacing="0" cellpadding="0" style="border-collapse:collapse;width:100%;margin:6px 0 18px">' \
               f"<tr>{h}</tr>{b}</table>"

    def link(text, url):
        return f'<a href="{_e(url)}" style="color:#1f5fa8">{_e(text)}</a>' if url else _e(text)

    def pct(v, invert=False):
        color = "#c0392b" if (v or 0) > 0 else "#1e8449"
        return f'<span style="color:{color};font-weight:bold">{"▲" if (v or 0) > 0 else "▼"} {abs(v or 0):.1f}%</span>'

    def section(title, body):
        return f'<h3 style="color:#1f3a5f;margin:22px 0 4px;font-size:16px">{_e(title)}</h3>{body}'

    kpi = lambda n, label: (f'<td style="padding:10px;text-align:center;background:#f3f6fa;border-radius:6px">'
                            f'<div style="font-size:22px;font-weight:bold;color:#1f3a5f">{n}</div>'
                            f'<div style="font-size:12px;color:#555">{_e(label)}</div></td>')
    parts = [
        '<div dir="rtl" style="font-family:Arial,Helvetica,sans-serif;max-width:760px;margin:auto;color:#222">',
        f'<h2 style="color:#1f3a5f;margin-bottom:2px">דוח מחירים יומי – {_e(traklin_name)} מול המתחרים</h2>',
        f'<div style="color:#666;font-size:13px">{date_he}' +
        (f" · השוואה לריצה מתאריך {_fmt_date(result.prev_date)}" if result.prev_date else
         " · ריצה ראשונה – מחר כבר יופיעו עליות וירידות") + "</div>",
    ]
    if result.partial:
        parts.append('<p style="color:#c0392b">⚠ הריצה נעצרה לפני הסוף בגלל מגבלת זמן – חלק מהמוצרים לא נבדקו היום.</p>')
    parts.append('<table cellspacing="8" style="width:100%;margin:12px 0"><tr>' +
                 kpi(len(result.products), "מוצרים בטרקלין") + kpi(found_total, "השוואות שנמצאו") +
                 kpi(len(moves), "שינויי מחיר") + kpi(promos_total, "מבצעים פעילים") + "</tr></table>")

    if cheaper:
        parts.append(section("איפה המתחרים זולים יותר מטרקלין", table(
            ["מוצר", "טרקלין", "הכי זול", "פער"],
            [[link(r["product"].title, r["product"].url), _money(r["product"].price),
              link(f'{r["store"]} – {_money(r["offer"].price)}', r["offer"].url),
              f'<b>{_money(r["gap"])}</b> ({r["gap_pct"]:.0f}%)'] for r in cheaper])))
    if downs:
        parts.append(section("ירידות מחיר", table(
            ["חנות", "מוצר", "היה", "עכשיו", "שינוי"],
            [[_e(c.site_name), link(c.title, c.url), _money(c.old), _money(c.new), pct(c.pct)] for c in downs])))
    if ups:
        parts.append(section("עליות מחיר", table(
            ["חנות", "מוצר", "היה", "עכשיו", "שינוי"],
            [[_e(c.site_name), link(c.title, c.url), _money(c.old), _money(c.new), pct(c.pct)] for c in ups])))
    if new_promos:
        parts.append(section("מבצעים חדשים", table(
            ["חנות", "מוצר", "מחיר"],
            [[_e(c.site_name), link(c.title, c.url), _money(c.new)] for c in new_promos])))
    if not result.first_run and not moves and not new_promos:
        parts.append('<p style="color:#555">לא היו שינויי מחיר או מבצעים חדשים מאז הריצה הקודמת.</p>')

    parts.append(section("מצב האתרים", table(
        ["חנות", "מצב", "נמצאו", "זול מטרקלין", "יקר מטרקלין"],
        [[_e(s["name"]), '<span style="color:#1e8449">✓ תקין</span>' if s["ok"] else
          '<span style="color:#c0392b">⚠ לא נקרא היום</span>', s["found"], s["cheaper"], s["dearer"]]
         for s in stats])))
    parts.append('<p style="color:#666;font-size:12px">הדוח המלא – כל המוצרים, כל החנויות, מבצעים וקישורים – '
                 'נמצא בקובץ האקסל המצורף.</p></div>')
    return subject, "".join(parts)
